from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path('output')
OUT.mkdir(exist_ok=True)

NAVY = '17324D'
BLUE = '2F6690'
LIGHT_BLUE = 'EAF2F8'
LIGHT_GRAY = 'F3F5F7'
BORDER = 'D9D9D9'


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=BORDER, size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_text(cell, text, bold=False, color='000000', size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_doc(title, subtitle):
    d = Document()
    sec = d.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    normal = d.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 35, 35)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in [('Title', 22, NAVY), ('Heading 1', 15, NAVY), ('Heading 2', 11.5, BLUE)]:
        s = d.styles[style_name]
        s.font.name = 'Aptos Display' if style_name == 'Title' else 'Aptos'
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(12 if style_name != 'Title' else 0)
        s.paragraph_format.space_after = Pt(5)
    p = d.add_paragraph(style='Title')
    p.add_run(title)
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(subtitle)
    r.italic = True
    r.font.name = 'Aptos'
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    header = sec.header.paragraphs[0]
    header.text = 'AI Dance Coach | GSB S5576 Project'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string('6B7280')
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Cal Poly Fall AI Convening target date: October 22, 2026')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('6B7280')
    return d


def add_bullets(d, items, level=0):
    for item in items:
        p = d.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_table(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_text(c, h, bold=True, color='FFFFFF', size=9)
        shade(c, NAVY)
        set_cell_border(c)
        if widths:
            c.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=8.8)
            shade(cells[i], 'FFFFFF' if ridx % 2 == 0 else LIGHT_GRAY)
            set_cell_border(cells[i])
            if widths:
                cells[i].width = Inches(widths[i])
    d.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def add_callout_paragraph(d, lead, text):
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(lead + ' ')
    r.bold = True
    p.add_run(text)


def build_short():
    d = make_doc('AI Dance Coach Project Plan', 'Short version for project planning, class discussion, and the Fall AI Convening application')
    d.add_heading('Project goal', level=1)
    d.add_paragraph('Build and deploy a Python-based web application that analyzes a short ballet rehearsal video and gives clear, non-diagnostic feedback on selected technique proxies and group synchronization. The tool supports practice and teacher guidance; it does not replace a dance instructor.')
    d.add_heading('MVP scope', level=1)
    add_bullets(d, [
        'Upload one short rehearsal video and validate its format, length, and visibility.',
        'Extract body landmarks frame by frame with MediaPipe Pose.',
        'Calculate posture, arm placement, lower-body alignment, balance, and synchronization proxies.',
        'Show annotated frames, metric charts, confidence levels, limitations, and three prioritized suggestions.',
        'Delete uploaded videos after analysis or after a clearly stated short retention period.'
    ])
    add_callout_paragraph(d, 'Scope boundary.', 'Treat turnout as a camera-dependent 2D proxy. Do not present it as a definitive technique or medical assessment.')
    d.add_heading('Recommended technology', level=1)
    add_table(d, ['Need', 'Recommended tool', 'Purpose'], [
        ('Web app', 'Streamlit', 'Fast Python interface for upload and results'),
        ('Video', 'OpenCV and FFmpeg', 'Read, sample, and annotate video'),
        ('Pose', 'MediaPipe Pose Landmarker', 'Detect body landmarks'),
        ('Analysis', 'NumPy, pandas, SciPy', 'Geometry, smoothing, and time-series metrics'),
        ('Charts', 'Plotly', 'Show movement and timing over time'),
        ('Feedback', 'Rules first; optional LLM', 'Keep suggestions grounded in measured values'),
        ('Deployment', 'Streamlit Community Cloud', 'Shareable demonstration URL')
    ], [1.05, 1.7, 4.45])
    d.add_heading('Data needed', level=1)
    add_bullets(d, [
        'Self-recorded or consented 5-20 second clips with a fixed camera and full body visible.',
        'Video metadata: frame rate, resolution, duration, camera view, dancer count, and movement label.',
        'Pose landmarks with timestamp, person identifier, normalized coordinates, and visibility confidence.',
        'A small reference set and expert review of whether generated suggestions are reasonable.',
        'User-test results covering clarity, usefulness, trust, and failure cases.'
    ])
    d.add_page_break()
    d.add_heading('Deadlines', level=1)
    add_table(d, ['Date', 'Target', 'Deliverable'], [
        ('Sep 22', 'Application', 'Submit Fall AI Convening application'),
        ('Oct 3', 'Working MVP', 'Upload, pose extraction, one metric, results page'),
        ('Oct 12-14', 'User testing', 'Three or more users complete a test task'),
        ('Oct 19', 'Demo freeze', 'Poster, QR code, backup recording, stable demo'),
        ('Oct 22', 'Presentation', 'Present poster and demonstrate the application'),
        ('Course Week 12', 'Final product', 'Deployed and polished application'),
        ('Course Week 13', 'Final presentation', 'Explain and demonstrate the final product')
    ], [1.05, 1.45, 4.7])
    d.add_heading('Success criteria', level=1)
    add_bullets(d, [
        'At least 80% of test clips produce usable landmarks and a results page.',
        'The app reports confidence and clear errors instead of inventing feedback.',
        'At least three people test the workflow and one experienced dancer reviews outputs.',
        'The poster includes the problem, pipeline, evidence, limitations, and responsible-AI choices.'
    ])
    d.add_heading('Start here', level=1)
    d.add_paragraph('Record two or three short test videos, create the project skeleton, and build one complete vertical slice: upload video, extract pose landmarks, calculate one metric, display one chart, and generate one grounded suggestion. Add more metrics only after this path works.')
    d.save(OUT / 'AI_Dance_Coach_Short_Project_Plan.docx')


def build_outline():
    d = make_doc('AI Dance Coach Build Outline', 'Step-by-step execution plan from first test video to deployed demonstration')
    d.add_heading('How to use this outline', level=1)
    d.add_paragraph('Work in the order shown. Each stage has a concrete output and a stopping point. If a stage is not working, fix it before adding the next feature. A smaller reliable demo is more valuable than a large unfinished system.')
    d.add_heading('Stage 1 Define the supported use case', level=1)
    d.add_paragraph('Choose one camera setup and one movement context for the first version. For example: two dancers, front-facing camera, full body visible, and a 10-second phrase.')
    add_bullets(d, ['Write the user problem in two sentences.', 'List the five metrics the MVP will attempt.', 'Write what the app will not claim to measure.', 'Confirm the course calendar and event requirements.'])
    add_callout_paragraph(d, 'Output.', 'A one-page MVP brief with target user, supported video setup, metrics, limitations, and success criteria.')
    d.add_heading('Stage 2 Prepare safe test data', level=1)
    add_bullets(d, ['Record two or three self-created test clips.', 'Keep the camera fixed and the entire body in frame.', 'Create a metadata sheet with duration, frame rate, camera view, dancer count, and movement label.', 'Add consent and deletion language before collecting anyone else’s video.'])
    add_callout_paragraph(d, 'Output.', 'A small test folder and data dictionary. Keep raw videos out of Git and out of application logs.')
    d.add_heading('Stage 3 Create the project skeleton', level=1)
    add_bullets(d, ['Create app.py, requirements.txt, src/, tests/, data/, notebooks/, and assets/.', 'Install Streamlit, OpenCV, MediaPipe, NumPy, pandas, Plotly, and pytest.', 'Add a README with setup instructions.', 'Create a basic Streamlit page with an upload control.'])
    add_callout_paragraph(d, 'Output.', 'The app runs locally and accepts a video without analyzing it yet.')
    d.add_heading('Stage 4 Build the pose pipeline', level=1)
    add_bullets(d, ['Read video metadata and sample frames.', 'Run pose detection on each frame.', 'Store timestamped landmarks and visibility confidence.', 'Draw a skeleton on sample frames.', 'Handle low confidence and no-person cases with clear messages.'])
    add_callout_paragraph(d, 'Output.', 'One test clip produces a landmark table and an annotated frame.')
    d.add_heading('Stage 5 Build one complete vertical slice', level=1)
    d.add_paragraph('Before building every metric, connect the full workflow using posture as the first metric.')
    add_bullets(d, ['Normalize landmarks by body size.', 'Calculate torso lean or shoulder-hip alignment.', 'Plot the value across time.', 'Create one rule-based suggestion.', 'Display the metric, confidence, chart, and limitation together.'])
    add_callout_paragraph(d, 'Output.', 'A working one-video MVP that can be demonstrated to another person.')
    d.add_heading('Stage 6 Add the remaining metrics', level=1)
    add_table(d, ['Order', 'Metric', 'First implementation'], [
        ('1', 'Arm placement', 'Shoulder-elbow-wrist angles'),
        ('2', 'Lower-body alignment', 'Hip-knee-ankle relationships in the camera view'),
        ('3', 'Balance', 'Torso or hip sway during a stable phrase'),
        ('4', 'Synchronization', 'Normalized landmark similarity and timing offset for two dancers')
    ], [0.65, 1.8, 4.75])
    d.add_paragraph('Add a unit test for each metric using simple, known landmark coordinates. Record the camera sensitivity and expected failure cases.')
    d.add_heading('Stage 7 Add feedback and responsible AI controls', level=1)
    add_bullets(d, ['Use measured values and thresholds to create feedback templates.', 'Show confidence and a limitation beside every recommendation.', 'If an LLM is used, send structured metrics instead of raw video.', 'Require the model to stay within the measured findings.', 'Keep a no-LLM fallback so the app remains functional without an API key.'])
    add_callout_paragraph(d, 'Output.', 'Feedback that is understandable, traceable to a metric, and safe to explain in a poster presentation.')
    d.add_heading('Stage 8 Deploy early', level=1)
    add_bullets(d, ['Deploy the simplest stable version before adding polish.', 'Test the public URL from another device and network.', 'Measure approximate processing time for a typical clip.', 'Prepare a backup video, screenshots, and a short screen recording.'])
    add_callout_paragraph(d, 'Output.', 'A working demo URL and a backup presentation path.')
    d.add_heading('Stage 9 Test with users', level=1)
    add_bullets(d, ['Give every tester the same task and instructions.', 'Ask whether the feedback is clear, useful, and trustworthy.', 'Record failures caused by lighting, camera angle, occlusion, clothing, or tracking.', 'Ask an experienced dancer to review at least five outputs.', 'Prioritize fixes by impact on successful use of the main workflow.'])
    add_callout_paragraph(d, 'Output.', 'A short testing summary with findings, quotes or ratings, failure cases, and a revision list.')
    d.add_heading('Stage 10 Prepare the poster', level=1)
    add_bullets(d, ['Use the working title AI Dance Coach: Pose-Based Feedback for Ballet Practice and Group Synchronization.', 'Show the problem, pipeline, annotated frame, metric chart, and testing evidence.', 'State that measurements are camera-dependent proxies.', 'Include privacy, consent, deletion, and human-in-the-loop choices.', 'Add a QR code to the demo or recorded walkthrough.'])
    add_callout_paragraph(d, 'Output.', 'A poster that a visitor can understand in under two minutes and a 60-90 second explanation.')
    d.add_heading('Stage 11 Freeze and present', level=1)
    add_bullets(d, ['Freeze the demo and poster by October 19.', 'Rehearse the live path and the backup path.', 'Prepare answers to: What problem does it solve? How does the AI work? What data does it use? How did you test it? What can it not do?', 'Present at the Fall AI Convening on October 22.', 'Continue improvements for the course Week 12 final product and Week 13 presentation.'])
    d.add_heading('First week checklist', level=1)
    add_table(d, ['Task', 'Done when'], [
        ('Confirm scope', 'One supported camera setup and one supported movement phrase are written down'),
        ('Record test clips', 'Two or three short videos exist and are safe to use'),
        ('Create project files', 'App launches locally and the README explains setup'),
        ('Build upload screen', 'A video can be selected and basic errors are handled'),
        ('Extract landmarks', 'An annotated frame and landmark table are produced'),
        ('Submit application', 'The Fall AI Convening form is submitted by September 22')
    ], [2.55, 4.85])
    d.add_paragraph('The best place to start today is the first-week checklist: record a test clip, create the app skeleton, and prove that one video can become one annotated pose frame.')
    d.save(OUT / 'AI_Dance_Coach_Build_Outline.docx')


if __name__ == '__main__':
    build_short()
    build_outline()
    print('Created:', OUT / 'AI_Dance_Coach_Short_Project_Plan.docx')
    print('Created:', OUT / 'AI_Dance_Coach_Build_Outline.docx')
