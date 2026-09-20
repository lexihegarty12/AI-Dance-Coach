from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "step two" / "AI Dance Coach Step Two Next Steps.docx"


def cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, color):
    from docx.oxml import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", color)
    tc_pr.append(shd)


def borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    bs = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        bs.append(e)
    tbl_pr.append(bs)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
for name in ("Title", "Heading 1", "Heading 2"):
    doc.styles[name].font.name = "Aptos"
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles[name].paragraph_format.keep_with_next = True

title = doc.add_paragraph(style="Title")
title.add_run("AI Dance Coach Step Two Next Steps")
sub = doc.add_paragraph()
sub.add_run("From posture prototype to visual technique feedback").italic = True
doc.add_paragraph("Step two will make the dashboard more useful to dancers by moving from a single torso-lean metric to position-aware visual feedback about lower-body alignment, hip orientation, and foot shape. The system should provide practice cues, not universal grades or medical judgments.")

doc.add_heading("Step two priorities", level=1)
for item in [
    "Render feedback directly over the dancer when pose confidence is high.",
    "Add a turnout and lower-body alignment proxy with camera limitations.",
    "Track knee-to-foot alignment during plié without treating knee-over-toe as automatically wrong.",
    "Estimate whether the pelvis remains square to the selected camera view.",
    "Detect possible sickling only when heel, ankle, and forefoot landmarks are visible.",
    "Use exercise-specific expectations for plié, passé, arabesque, penché, and pirouette phrases.",
]:
    bullet(doc, item)

doc.add_heading("Planned measurements", level=1)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, text in zip(table.rows[0].cells, ["Area", "First proxy", "Visual feedback", "Limitation"]):
    cell_text(c, text, bold=True, color="FFFFFF")
    shade(c, "334E68")
rows = [
    ("Turnout", "Projected foot orientation and hip-knee-ankle relationships", "Highlight lower legs and show a short cue", "A side view cannot establish true hip turnout"),
    ("Knee alignment", "Knee position relative to the second-toe direction during plié", "Draw a knee-to-foot direction line", "Knee-over-toe is not a universal injury rule"),
    ("Square hips", "Pelvis line and left/right hip relationships", "Show the hip line and indicate rotation", "Perspective can resemble hip rotation"),
    ("Sickled foot", "Heel, ankle, and forefoot line", "Outline the foot only when visible", "Foot landmarks are noisy"),
    ("Arm placement", "Shoulder-elbow-wrist angles and hand height", "Color the arm chain and show the position label", "Placement depends on phrase and style"),
]
for row in rows:
    cells = table.add_row().cells
    for c, text in zip(cells, row):
        cell_text(c, text)
borders(table)

doc.add_heading("Implementation sequence", level=1)
sequence = [
    ("1. Save landmark tracks", "Retain the relevant hip, knee, ankle, heel, foot-index, shoulder, elbow, and wrist coordinates plus visibility confidence for every frame."),
    ("2. Add exercise context", "Keep the user-provided clip labels and add a position or phrase field. Thresholds should be exercise-specific."),
    ("3. Build knee tracking", "Start with first-position and fifth-position pliés. Flag sustained, high-confidence deviations rather than single noisy frames."),
    ("4. Add hip alignment", "Estimate a pelvis-line cue and distinguish camera-view rotation from a claim about anatomical position."),
    ("5. Add turnout proxy", "Use front or three-quarter reference clips and report confidence and camera sensitivity with every result."),
    ("6. Add sickling proxy", "Use foot orientation during tendu, relevé, or extension only when foot landmarks are visible."),
    ("7. Render visual feedback", "Overlay the skeleton, relevant alignment line, metric, and one concise practice cue on the video."),
    ("8. Validate", "Have an experienced dancer or instructor label at least five examples per measurement as reasonable, unclear, or wrong."),
]
for heading, body in sequence:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run(heading).bold = True
    doc.add_paragraph(body)

doc.add_heading("Reference video requirements", level=1)
for item in [
    "Use self-recorded, consented, licensed, or public-domain reference clips.",
    "Record full body, feet, and arms with a stationary camera.",
    "Capture front and three-quarter views for turnout, hip, knee, and foot measures.",
    "Use side views mainly for torso and balance observations.",
    "Label each clip with exercise, camera view, and movement phase.",
    "Ask an experienced dancer or instructor to review the first feedback outputs.",
]:
    bullet(doc, item)

doc.add_heading("Definition of done", level=1)
doc.add_paragraph("Step two is complete when a dancer can select a labeled clip and see a visual, time-linked cue for the selected measurement without opening a CSV file.")
for item in [
    "The feedback video plays in the dashboard.",
    "The overlay shows the relevant body line or joint relationship.",
    "Low-confidence and unsuitable camera views are visibly disclosed.",
    "The app distinguishes an exercise-specific practice cue from a universal technique judgment.",
    "At least one experienced dancer or instructor has reviewed the outputs.",
]:
    bullet(doc, item)

doc.add_heading("Recommended immediate task", level=1)
doc.add_paragraph("Begin with plies_first, plies_fifth, and passe_balance. Implement knee tracking and the visual feedback renderer before expanding to developpé and pirouette phrases.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
